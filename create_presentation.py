from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_presentation():
    # Create a presentation object
    prs = Presentation()
    
    # Define slide layouts
    title_slide_layout = prs.slide_layouts[0]
    title_and_content_layout = prs.slide_layouts[1]
    
    # Helper function to add a slide with title and content
    def add_slide(title, content_list):
        slide = prs.slides.add_slide(title_and_content_layout)
        title_shape = slide.shapes.title
        title_shape.text = title
        
        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame
        tf.clear()
        
        for i, point in enumerate(content_list):
            p = tf.add_paragraph()
            p.text = point
            p.font.size = Pt(20)
            if i > 0:
                p.space_before = Pt(10)
        return slide

    # Slide 1: Title Slide
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "NEXUSCOMMERCE"
    title.text_frame.paragraphs[0].font.size = Pt(44)
    title.text_frame.paragraphs[0].font.bold = True
    
    subtitle.text = "An AI-Driven E-Commerce Intelligence and Decision Support System\n\nBy:\nMuhammad Abdullah (CIIT/SP23-BAI-027/ISB)\nLayiba Aslam (CIIT/SP23-BAI-054/ISB)\n\nSupervisor: Dr. M. Manzoor Illahi Tamimi\nCOMSATS University, Islamabad"

    # Slide 2: Executive Summary
    add_slide("Executive Summary", [
        "NexusCommerce is an AI-driven e-commerce decision support system.",
        "Addresses operational challenges faced by small/medium online sellers.",
        "Replaces fragmented tools and manual spreadsheet analysis with an intelligent platform.",
        "Unifies 6 critical business functions into a single system.",
        "Transforms operations by replacing guesswork with data-driven intelligence and manual processes with AI automation."
    ])

    # Slide 3: Problem Statement
    add_slide("Problem Statement", [
        "E-commerce sellers lack access to affordable, integrated, intelligent tools.",
        "Current reliance on disconnected tools (marketplaces, spreadsheets, manual tracking).",
        "Results in data silos, delayed decision-making, and missed opportunities.",
        "Key Challenges:",
        "- Inaccurate demand prediction leading to stockouts or excess inventory.",
        "- Reactive pricing decisions without true profitability insight.",
        "- Manual, time-consuming supplier ordering and reporting.",
        "- Lack of a conversational interface to query business data."
    ])

    # Slide 4: Proposed Solution
    add_slide("Proposed Solution: NexusCommerce", [
        "An end-to-end intelligent platform transforming raw sales data into actionable insights.",
        "Provides:",
        "- Automated pipeline for data validation, cleaning, and anomaly detection.",
        "- XGBoost machine learning for demand forecasting with 95% confidence intervals.",
        "- ML-based risk scoring for proactive inventory management.",
        "- Agentic AI Assistant for data querying and business action execution (e.g., ordering).",
        "- Profitability-focused pricing intelligence.",
        "- Interactive dashboards with real-time filtering and business health scoring."
    ])

    # Slide 5: Key Objectives
    add_slide("Key Objectives", [
        "Improve Demand Prediction: Achieve at least 85% accuracy using XGBoost with seller-specific fine-tuning.",
        "Optimize Inventory: Reduce stockouts via ML-based risk scoring with 90% prediction accuracy.",
        "Enhance Profitability: Improve margins by at least 10% through ML-based pricing recommendations.",
        "Automate Operations: Enable natural language execution of routine actions via Agentic AI.",
        "Ensure Data Quality: Automated quality scoring, deduplication, and version history.",
        "Unified Platform: Consolidate data, forecasting, inventory, pricing, and AI into one system."
    ])

    # Slide 6: System Scope & Modules
    add_slide("System Scope & Core Modules", [
        "NexusCommerce is built on a modular three-tier architecture.",
        "Core Modules:",
        "1. Data Integration & Management",
        "2. Demand Forecasting (Core AI)",
        "3. Inventory Intelligence & Supplier Management",
        "4. AI Conversational Assistant (Agentic AI)",
        "5. Dashboard & Reporting Interface",
        "6. Pricing Intelligence & Profitability"
    ])

    # Slide 7: System Architecture
    add_slide("System Architecture", [
        "Frontend Layer: Next.js + React + Tailwind CSS (Interactive Dashboards, Chat Interface)",
        "Backend API Layer: Node.js + Express.js (Business Logic, Auth, File Processing)",
        "AI/ML Processing Layer: Python FastAPI + XGBoost (Forecasting, Fine-tuning, Risk Scoring)",
        "Database Layer: PostgreSQL (Centralized storage for 17 interconnected tables)",
        "Agentic AI Layer: RAG + MCP + External LLMs",
        "External Integrations: n8n workflow automation for execution (e.g., placing orders, sending reports)",
        "Security: JWT Auth, bcryptjs, Server-side Seller Isolation, RBAC"
    ])

    # Slide 8: Module 1 - Data Integration
    add_slide("Module 1: Data Integration & Management", [
        "Accepts manual uploads (CSV, JSON, Excel).",
        "Automated Pipeline:",
        "- Per-row validation with severity-based anomaly classification.",
        "- Data cleaning (median-based imputation, IQR outlier detection).",
        "- Automatic deduplication using PostgreSQL unique constraints.",
        "Generates quality score (0-100) after every upload.",
        "Maintains version history with safe rollback capability.",
        "Tracks data freshness to ensure analytics use current data."
    ])

    # Slide 9: Module 2 - Demand Forecasting
    add_slide("Module 2: Demand Forecasting", [
        "Powered by XGBoost Regressor for recursive day-by-day prediction.",
        "Product-level forecasts for 7, 14, and 30-day horizons.",
        "Provides 95% confidence intervals and feature importance explainability.",
        "3-Tier Model Priority:",
        "1. Seller-specific fine-tuned model",
        "2. Category base model",
        "3. Fallback model",
        "Supports one-click warm-start fine-tuning to personalize models.",
        "Performance evaluated via MAE, RMSE, R², and Custom Accuracy Score."
    ])

    # Slide 10: Module 3 - Inventory Intelligence
    add_slide("Module 3: Inventory Intelligence & Supplier Management", [
        "Goes beyond static thresholds using ML-based risk scoring.",
        "Risk factors: Current stock, forecasted demand, sales velocity, supplier lead times.",
        "Risk levels: Critical, High, Medium, Low.",
        "Calculates exact order deadlines and recommends specific restock quantities.",
        "Quantifies financial impact (estimated lost revenue vs. restock costs).",
        "Manages supplier profiles and compares options.",
        "Provides an overall 'Inventory Health Score'."
    ])

    # Slide 11: Module 4 - AI Conversational Assistant
    add_slide("Module 4: AI Conversational Assistant (Agentic AI)", [
        "Operates in two powerful modes:",
        "1. Query Mode:",
        "- Uses Retrieval Augmented Generation (RAG) and Model Context Protocol (MCP).",
        "- Answers business questions grounded in real PostgreSQL data.",
        "2. Action Mode:",
        "- Enables execution of real operations via n8n workflows.",
        "- Example: \"Order 200 units of Product A\".",
        "- Generates confirmation cards before execution.",
        "- Double confirmation required for high-value actions.",
        "- Maintains complete immutable audit trail of AI actions."
    ])

    # Slide 12: Module 5 & 6 - Dashboard & Pricing
    add_slide("Dashboard & Pricing Intelligence", [
        "Module 5: Dashboard & Reporting",
        "- Interactive, real-time filtering without page reloads.",
        "- Displays Business Health Score, KPI cards, anomaly alerts.",
        "- Role-based views (Sellers, Analysts, Admins).",
        "Module 6: Pricing Intelligence",
        "- Tracks cost prices to compute actual profit margins.",
        "- ML-based optimal price recommendations maximizing profitability.",
        "- Explains reasoning and predicts revenue impact.",
        "- Ensures a minimum 15% margin and never prices below cost."
    ])

    # Slide 13: Project Contributions (Novelty)
    add_slide("Project Contributions", [
        "Seller-Personalized Demand Forecasting: Warm-start fine-tuning over generic models.",
        "Agentic AI with Real Action Execution: Not just answering questions, but safely executing business actions via n8n.",
        "Financial Impact Quantification: Converting stock risks into concrete financial values (lost revenue).",
        "Profitability-First Pricing: Shifting focus from raw revenue to actual profit margins.",
        "Comprehensive Data Pipeline: Institutional-grade data quality, versioning, and freshness tracking for small sellers."
    ])

    # Slide 14: Conclusion
    add_slide("Conclusion", [
        "NexusCommerce provides a holistic, AI-powered ecosystem.",
        "Replaces guesswork and disjointed systems with a unified platform.",
        "Empowers SMEs to make data-driven decisions on demand, inventory, and pricing.",
        "Innovates by bridging analytical insights with conversational business automation.",
        "Sellers remain in full control of every business decision with explicit confirmation workflows."
    ])

    # Slide 15: Q&A
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "Thank You!"
    subtitle.text = "Any Questions?"

    # Save presentation
    prs.save("NexusCommerce_Presentation.pptx")
    print("Presentation created successfully as 'NexusCommerce_Presentation.pptx'")

if __name__ == "__main__":
    create_presentation()
